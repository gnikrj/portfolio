"""Build the Bending Spoons acquisition-integration model (Excel).

Every hardcoded number lives on the Inputs tab with its source; every other
tab is formulas off Inputs. Public sources only (Q2 2026 earnings release,
deal press releases). Run from this scripts/ directory:

    python 01_build_model.py
    python <xlsx skill>/scripts/recalc.py ../data/bending_spoons_integration_model.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule

OUT = Path(__file__).resolve().parent.parent / "data" / "bending_spoons_integration_model.xlsx"

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HDR = Font(name=FONT, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
YELLOW = PatternFill("solid", fgColor="FFFF00")
GREY = PatternFill("solid", fgColor="F2F2F2")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
WRAP = Alignment(wrap_text=True, vertical="top")

USD = '$#,##0;($#,##0);"-"'
USD1 = '$#,##0.0;($#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.00"x"'
DAYS = '0" days"'

REL = "Q2 2026 earnings release (Aug 13, 2026)"
REL_URL = "https://bendingspoons.com/documents/financials/2026/20260813%20Bending%20Spoons%E2%80%94Q2%202026%20Results.pdf"
MIRO_URL = "https://investors.bendingspoons.com/newsroom/bending-spoons-agrees-to-acquire-miro"
AIRTABLE_URL = "https://investors.bendingspoons.com/newsroom/bending-spoons-completes-acquisition-of-airtable"
GOOGLY_URL = "https://thegoogly.substack.com/p/bending-spoons-the-comeback-story"
MM_URL = "https://www.mostlymetrics.com/p/bending-spoons-ipo-s1-breakdown"

wb = Workbook()


def style_ws(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False


def header(ws, row, labels):
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill, c.alignment = HDR, HDR_FILL, Alignment(wrap_text=True, vertical="center")


def put(ws, ref, value, font=BLACK, fmt=None, fill=None, comment=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if comment:
        c.comment = Comment(comment, "Jeff King")
    return c


# ---------------------------------------------------------------- Cover
ws = wb.active
ws.title = "Cover"
style_ws(ws, {"A": 3, "B": 110})
lines = [
    ("Bending Spoons: acquisition-integration model", TITLE),
    ("Outside-in, public sources only. As of Oct 12, 2026. Prepared by Jeff King.", BLACK),
    ("", BLACK),
    ("Question: how much does integration speed matter to Bending Spoons' economics, and what should finance own in each integration?", BOLD),
    ("", BLACK),
    ("Tabs", BOLD),
    ("Inputs: every hardcoded number, with its source. Change blue cells here; everything else recalculates.", BLACK),
    ("Leverage: net-debt walk from Q2 2026 through the IPO, Airtable and Miro, vs. the 4.0x covenant. Shows why documented cost savings matter.", BLACK),
    ("Miro_Payback: price / post-integration operating income across revenue-retention and margin cases; cost of each month of delay.", BLACK),
    ("Working_Capital: receivables and DSO check across acquisitions (order-to-cash lens).", BLACK),
    ("Scorecard: one-page integration scorecard template per acquisition (example row = Miro, illustrative targets).", BLACK),
    ("Lookback: 12-month underwritten-vs-actual template that feeds the next underwriting.", BLACK),
    ("Sources: links to every page used.", BLACK),
    ("", BLACK),
    ("Legend", BOLD),
    ("Blue text = hardcoded input from a source  |  Black = formula  |  Green = link to another tab  |  Yellow fill = my assumption (change it)", BLACK),
    ("", BLACK),
    ("Caveats", BOLD),
    ("Reported leverage uses pro forma adjusted EBITDA that includes achieved AND certain expected cost savings (release footnote 3). TTM EBITDA here is implied from net debt / 2.4x, so it carries rounding error of about +/-2%.", BLACK),
    ("The 4.0x covenant figure comes from a secondary summary of the IPO prospectus, not the filing itself. Verify before quoting.", BLACK),
    ("Airtable and Miro EBITDA are not disclosed; their contributions are assumptions. Nothing here is investment advice.", BLACK),
]
for i, (text, font) in enumerate(lines, start=1):
    c = ws.cell(row=i, column=2, value=text)
    c.font = font
    c.alignment = WRAP

# ---------------------------------------------------------------- Inputs
ws = wb.create_sheet("Inputs")
style_ws(ws, {"A": 52, "B": 14, "C": 10, "D": 46, "E": 60})
put(ws, "A1", "Inputs ($ in millions unless noted)", TITLE)
header(ws, 3, ["Item", "Value", "Unit", "Source", "Note"])

INPUTS = [
    # key, label, value, unit, fmt, source, note, assumption?
    ("sec", "Bending Spoons, Q2 2026 (quarter ended Jun 30, 2026)"),
    ("q2_rev", "Q2 2026 revenue", 704.155, "$mm", USD1, REL, "Income statement", False),
    ("q2_aoi", "Q2 2026 Adjusted Operating Income", 381.149, "$mm", USD1, REL, "Non-GAAP reconciliation", False),
    ("q2_int", "Q2 2026 interest expense", 108.970, "$mm", USD1, REL, "Income statement", False),
    ("q2_reorg", "Q2 2026 reorganization-related expense", 50.799, "$mm", USD1, REL, "Mostly separation packages at AOL, Eventbrite, Tractive, Vimeo", False),
    ("org_g", "Q2 2026 organic revenue growth", 0.03, "%", PCT, REL, "Headline growth was 126%", False),
    ("nd_q2", "Net debt, Jun 30, 2026", 4090, "$mm", USD, REL, "Long-term debt $4.88B less cash $793M", False),
    ("lev_q2", "Reported leverage ratio, Jun 30, 2026", 2.4, "x", MULT, REL, "Net debt / TTM pro forma adj. EBITDA incl. achieved + certain expected savings (footnote 3)", False),
    ("cash_q2", "Cash and cash equivalents, Jun 30, 2026", 792.950, "$mm", USD1, REL, "Balance sheet", False),
    ("debt_12m", "Debt maturing within 12 months", 794, "$mm", USD, REL, "Long-term debt, current", False),
    ("rcf_undrawn", "Undrawn revolver capacity", 1280, "$mm", USD, REL, "Of $1.58B total", False),
    ("ocf_h1", "H1 2026 net cash from operating activities", 254.240, "$mm", USD1, REL, "After $161M interest paid", False),
    ("capex_h1", "H1 2026 capex (PP&E)", 3.969, "$mm", USD1, REL, "Cash flow statement", False),
    ("ar_dec", "Accounts receivable, Dec 31, 2025", 144.593, "$mm", USD1, REL, "Balance sheet", False),
    ("ar_jun", "Accounts receivable, Jun 30, 2026", 278.767, "$mm", USD1, REL, "Balance sheet", False),
    ("rev_h1_25", "H1 2025 revenue", 570.046, "$mm", USD1, REL, "Six months ended Jun 30, 2025", False),
    ("rev_fy25", "FY2025 revenue", 1306, "$mm", USD, MM_URL, "Secondary summary of prospectus", False),
    ("fy_aoi_lo", "FY2026 Adjusted Operating Income guidance, low", 1460, "$mm", USD, REL, "Owned businesses as of Aug 12, 2026", False),
    ("fy_aoi_hi", "FY2026 Adjusted Operating Income guidance, high", 1510, "$mm", USD, REL, "Excludes Airtable", False),
    ("sec", "Post-quarter events"),
    ("ipo", "IPO net proceeds (July 2026)", 1100, "$mm", USD, REL, "39.6M primary shares at $29; assumed held as cash, reducing net debt", False),
    ("airtable_ev", "Airtable enterprise value (closed Sep 4, 2026)", 1290, "$mm", USD, REL, "All cash; EV basis so net debt rises by EV", False),
    ("miro_ev", "Miro enterprise value (signed Sep 10, 2026; close Q4)", 1355, "$mm", USD, MIRO_URL, "Equity value ~$1.79B incl. Miro net cash", False),
    ("miro_roll", "Miro sellers reinvesting in new BSP equity", 295, "$mm", USD, MIRO_URL, "Reduces cash funding need", False),
    ("miro_arr", "Miro ARR", 600, "$mm", USD, MIRO_URL, "'About $600 million'; ~90% business/enterprise", False),
    ("cov", "Leverage covenant (net debt / adj. EBITDA)", 4.0, "x", MULT, GOOGLY_URL, "Secondary summary of prospectus; verify in filing", False),
    ("rate", "Cost of incremental debt", 0.0943, "%", PCT, GOOGLY_URL, "All-in rate on largest USD term loan; euro TLAs likely cheaper", False),
    ("sec", "My assumptions (yellow; change these)"),
    ("fcf_h2", "H2 2026 free cash flow retained (debt paydown)", 300, "$mm", USD, "Assumption", "H1 OCF less capex was ~$250M; H2 has a fuller AOI base", True),
    ("airtable_ebitda", "Airtable pro forma EBITDA credited at year-end", 0, "$mm", USD, "Assumption", "Not disclosed. Base case 0 = no credit until savings are documented", True),
    ("miro_ebitda", "Miro pro forma EBITDA credited at year-end", 0, "$mm", USD, "Assumption", "Not disclosed. Base case 0", True),
    ("miro_ret", "Miro revenue retained after repricing (base)", 0.92, "%", PCT, "Assumption", "BSP group NRR was 91-95% in 2023-25 (secondary)", True),
    ("miro_margin", "Miro post-integration operating margin (base)", 0.45, "%", PCT, "Assumption", "Group adj. operating margin was 54% in Q2 2026", True),
]

R = {}  # key -> absolute ref on Inputs, e.g. Inputs!$B$5
row = 4
for item in INPUTS:
    if item[0] == "sec":
        ws.cell(row=row, column=1, value=item[1]).font = BOLD
        for col in range(1, 6):
            ws.cell(row=row, column=col).fill = GREY
        row += 1
        continue
    key, label, val, unit, fmt, src, note, assum = item
    ws.cell(row=row, column=1, value=label).font = BLACK
    c = ws.cell(row=row, column=2, value=val)
    c.font, c.number_format = BLUE, fmt
    if assum:
        c.fill = YELLOW
    ws.cell(row=row, column=3, value=unit).font = BLACK
    s = ws.cell(row=row, column=4, value=src)
    s.font = BLACK
    if src.startswith("http"):
        s.hyperlink = src
    elif src == REL:
        s.hyperlink = REL_URL
    ws.cell(row=row, column=5, value=note).font = BLACK
    ws.cell(row=row, column=5).alignment = WRAP
    for col in range(1, 6):
        ws.cell(row=row, column=col).border = THIN
    R[key] = f"Inputs!$B${row}"
    row += 1
ws.freeze_panes = "A4"


def g(ws, ref, key, fmt):
    """Green link to an Inputs cell."""
    return put(ws, ref, f"={R[key]}", GREEN, fmt)


# ---------------------------------------------------------------- Leverage
ws = wb.create_sheet("Leverage")
style_ws(ws, {"A": 58, "B": 16, "C": 70})
put(ws, "A1", "Net-debt and leverage walk, Jun 30, 2026 to pro forma year-end 2026 ($mm)", TITLE)
put(ws, "A2", "Point: with zero EBITDA credit for Airtable and Miro, leverage stays under the covenant, but the headroom left for the next deal depends on how fast savings are achieved and documented.", BLACK).alignment = WRAP
ws.row_dimensions[2].height = 30
header(ws, 4, ["Line", "Value", "How it's built"])
rows = [
    ("Net debt, Jun 30, 2026", f"={R['nd_q2']}", USD, GREEN, "Reported"),
    ("Less: IPO net proceeds", f"=-{R['ipo']}", USD, GREEN, "Assumes proceeds held as cash until deployed"),
    ("Less: H2 2026 free cash flow retained", f"=-{R['fcf_h2']}", USD, GREEN, "Assumption"),
    ("Plus: Airtable (EV)", f"={R['airtable_ev']}", USD, GREEN, "All cash, closed Sep 4, 2026"),
    ("Plus: Miro EV less seller rollover", f"={R['miro_ev']}-{R['miro_roll']}", USD, BLACK, "Expected Q4 2026 close"),
    ("Pro forma net debt, year-end 2026", "=SUM(B5:B9)", USD, BLACK, "Excludes Tractive $115M deferred consideration (May 2027) and deal fees"),
    ("", None, None, BLACK, ""),
    ("Implied TTM adj. EBITDA, Jun 30, 2026", f"={R['nd_q2']}/{R['lev_q2']}", USD, BLACK, "Net debt / reported 2.4x; held flat (conservative: ignores H2 growth in TTM)"),
    ("Plus: Airtable EBITDA credited", f"={R['airtable_ebitda']}", USD, GREEN, "Assumption"),
    ("Plus: Miro EBITDA credited", f"={R['miro_ebitda']}", USD, GREEN, "Assumption"),
    ("Pro forma adj. EBITDA", "=SUM(B12:B14)", USD, BLACK, ""),
    ("", None, None, BLACK, ""),
    ("Pro forma leverage", "=IFERROR(B10/B15,0)", MULT, BLACK, "Net debt / EBITDA"),
    ("Covenant", f"={R['cov']}", MULT, GREEN, "Secondary source; verify"),
    ("Headroom to covenant", "=B18-B17", MULT, BLACK, ""),
    ("Debt capacity before covenant ($mm)", "=B18*B15-B10", USD, BLACK, "Extra net debt before hitting 4.0x"),
    ("", None, None, BLACK, ""),
    ("Interest on Miro funding, annual ($mm)", f"=({R['miro_ev']}-{R['miro_roll']})*{R['rate']}", USD, BLACK, "Carry on the incremental debt at the cited 9.43% rate; runs from day one regardless of integration pace"),
    ("Q2 2026 interest as % of Adjusted Operating Income", f"={R['q2_int']}/{R['q2_aoi']}", PCT, BLACK, "Reported Q2 figures"),
]
for i, (lab, f, fmt, font, how) in enumerate(rows, start=5):
    ws.cell(row=i, column=1, value=lab).font = BOLD if lab.startswith(("Pro forma", "Headroom", "Debt capacity")) else BLACK
    if f is not None:
        put(ws, f"B{i}", f, font, fmt)
    ws.cell(row=i, column=3, value=how).font = BLACK
    ws.cell(row=i, column=3).alignment = WRAP

# EBITDA-credit sensitivity: combined Airtable+Miro credited EBITDA -> leverage
put(ws, "A26", "Sensitivity: combined Airtable + Miro EBITDA credited ($mm) vs. pro forma leverage and capacity", BOLD)
header(ws, 27, ["Combined EBITDA credited ($mm)", "Leverage", "Debt capacity to 4.0x ($mm)"])
for j, credit in enumerate([0, 100, 200, 300, 400, 500]):
    r = 28 + j
    put(ws, f"A{r}", credit, BLUE, USD)
    put(ws, f"B{r}", f"=IFERROR($B$10/($B$12+A{r}),0)", BLACK, MULT)
    put(ws, f"C{r}", f"=$B$18*($B$12+A{r})-$B$10", BLACK, USD)
put(ws, "A35", "Reading: each $100M of documented savings credited adds about $400M of acquisition capacity at 4.0x. That is why finance's savings evidence is part of the deal engine, not bookkeeping.", BLACK).alignment = WRAP
ws.row_dimensions[35].height = 30

# ---------------------------------------------------------------- Miro payback
ws = wb.create_sheet("Miro_Payback")
style_ws(ws, {"A": 46, "B": 14, "C": 14, "D": 14, "E": 14, "F": 40})
put(ws, "A1", "Miro: price / post-integration operating income", TITLE)
put(ws, "A2", "Multiple = EV / (ARR x revenue retained x operating margin). Lower is better; ~5x = 20% annual pre-tax return on price.", BLACK)
put(ws, "A4", "Miro EV ($mm)", BOLD); g(ws, "B4", "miro_ev", USD)
put(ws, "A5", "Miro ARR ($mm)", BOLD); g(ws, "B5", "miro_arr", USD)
put(ws, "A6", "EV / ARR", BOLD); put(ws, "B6", "=B4/B5", BLACK, MULT)

put(ws, "A8", "Grid: rows = operating margin, columns = revenue retained", BOLD)
put(ws, "A9", "Margin \\ Retained", HDR, fill=HDR_FILL)
for j, ret in enumerate([0.85, 0.92, 1.00]):
    col = "BCD"[j]
    put(ws, f"{col}9", ret, Font(name=FONT, bold=True, color="0000FF"), PCT, fill=GREY)
for i, m in enumerate([0.35, 0.45, 0.55]):
    r = 10 + i
    put(ws, f"A{r}", m, BLUE, PCT, fill=GREY)
    for j in range(3):
        col = "BCD"[j]
        put(ws, f"{col}{r}", f"=IFERROR($B$4/($B$5*{col}$9*$A{r}),0)", BLACK, MULT)

put(ws, "A15", "Base case (from Inputs)", BOLD)
put(ws, "A16", "Revenue retained"); g(ws, "B16", "miro_ret", PCT)
put(ws, "A17", "Operating margin"); g(ws, "B17", "miro_margin", PCT)
put(ws, "A18", "Run-rate operating income ($mm)"); put(ws, "B18", "=B5*B16*B17", BLACK, USD)
put(ws, "A19", "EV / run-rate operating income"); put(ws, "B19", "=IFERROR(B4/B18,0)", BLACK, MULT)
put(ws, "A20", "Pre-tax yield on EV"); put(ws, "B20", "=IFERROR(B18/B4,0)", BLACK, PCT)
put(ws, "A21", "Annual interest on Miro funding ($mm)"); put(ws, "B21", "=Leverage!B22", GREEN, USD)
put(ws, "A22", "Operating income / interest (coverage)"); put(ws, "B22", "=IFERROR(B18/B21,0)", BLACK, MULT)

put(ws, "A24", "Cost of delay", BOLD)
put(ws, "A25", "Operating income lost per month of delay ($mm)"); put(ws, "B25", "=B18/12", BLACK, USD1)
put(ws, "A26", "Plus interest carried per month ($mm)"); put(ws, "B26", "=B21/12", BLACK, USD1)
put(ws, "A27", "Acquisition capacity unlocked once savings are credited ($mm)"); put(ws, "B27", f"=B18*{R['cov']}", BLACK, USD)
put(ws, "F27", "Run-rate operating income x 4.0x covenant; used as a proxy for EBITDA", BLACK).alignment = WRAP
put(ws, "A29", "Risk note: Brightcove's revenue fell after acquisition (secondary sources). Retention, not margin, is the variable to watch; the look-back tab tracks it.", BLACK).alignment = WRAP
ws.row_dimensions[29].height = 30

# ---------------------------------------------------------------- Working capital
ws = wb.create_sheet("Working_Capital")
style_ws(ws, {"A": 46, "B": 16, "C": 16, "D": 50})
put(ws, "A1", "Receivables check: did AR grow faster than revenue as acquisitions landed?", TITLE)
header(ws, 3, ["Metric", "Dec 31, 2025", "Jun 30, 2026", "Note"])
put(ws, "A4", "Accounts receivable ($mm)"); g(ws, "B4", "ar_dec", USD1); g(ws, "C4", "ar_jun", USD1)
put(ws, "A5", "Quarter revenue ($mm)")
put(ws, "B5", f"=({R['rev_fy25']}-{R['rev_h1_25']})/2", BLACK, USD1)
put(ws, "D5", "Q4 2025 approximated as half of H2 2025 (not disclosed separately)").alignment = WRAP
g(ws, "C5", "q2_rev", USD1)
put(ws, "A6", "Days in quarter", BLACK); put(ws, "B6", 92, BLUE); put(ws, "C6", 91, BLUE)
put(ws, "A7", "DSO (AR / quarter revenue x days)", BOLD)
put(ws, "B7", "=IFERROR(B4/B5*B6,0)", BLACK, DAYS); put(ws, "C7", "=IFERROR(C4/C5*C6,0)", BLACK, DAYS)
put(ws, "A8", "AR growth"); put(ws, "C8", "=C4/B4-1", BLACK, PCT)
put(ws, "A9", "Revenue growth (quarter vs quarter)"); put(ws, "C9", "=C5/B5-1", BLACK, PCT)
put(ws, "A11", "Reading: AR roughly doubled, but DSO is about flat, so collections kept pace with acquired revenue. The mix is shifting, though: advertising (AOL) and invoiced B2B (Vimeo, Airtable, Miro) carry real credit and collection risk that app-store revenue doesn't. That argues for bringing each target's AR aging, credit memos, and bad-debt reserve into the Day 1-30 integration checklist.", BLACK).alignment = WRAP
ws.merge_cells("A11:D11")
ws.row_dimensions[11].height = 62

# ---------------------------------------------------------------- Scorecard
ws = wb.create_sheet("Scorecard")
style_ws(ws, {"A": 34, "B": 14, "C": 14, "D": 14, "E": 12, "F": 46})
put(ws, "A1", "Integration scorecard (one page per acquisition)", TITLE)
put(ws, "A2", "Fill the yellow cells. Example below = Miro, with illustrative targets I would propose; actuals are blank until close.", BLACK)
put(ws, "A3", "Acquisition", BOLD); put(ws, "B3", "Miro", BLUE, fill=YELLOW)
put(ws, "A4", "Close date", BOLD); put(ws, "B4", "Q4 2026 (expected)", BLUE, fill=YELLOW)
put(ws, "A5", "Owner (finance)", BOLD); put(ws, "B5", "Jeff King", BLUE, fill=YELLOW)
header(ws, 7, ["Metric", "Target", "Actual", "Variance", "Status", "Why it matters"])
SC = [
    ("Days to cash control (payouts on Juno)", 30, DAYS, "lower", "Cash is the first thing that can leak; 75% of revenue is electronic payments"),
    ("Days to first consolidated close", 60, DAYS, "lower", "Fixed external deadlines; the disclosed material weakness covers acquired-company integration"),
    ("Access removed same day as exits (%)", 1.0, PCT, "higher", "Segregation of duties breaks first in deep reorganizations (Galf)"),
    ("Run-rate savings achieved vs. plan (%)", 1.0, PCT, "higher", "Achieved savings are credited in reported leverage"),
    ("Net revenue retention vs. pre-deal (pp)", 0.0, PCT, "higher", "Brightcove lesson: margin up, revenue down"),
    ("Reorganization spend vs. budget (%)", 1.0, PCT, "lower", "Add-backs must reconcile to a cost center to survive audit"),
    ("Target systems retired (of total)", 1.0, PCT, "higher", "Simplification: every legacy system kept is a control to test"),
]
for i, (metric, tgt, fmt, better, why) in enumerate(SC, start=8):
    put(ws, f"A{i}", metric)
    put(ws, f"B{i}", tgt, BLUE, fmt, fill=YELLOW)
    put(ws, f"C{i}", None, BLUE, fmt, fill=YELLOW)
    put(ws, f"D{i}", f'=IF(C{i}="","",C{i}-B{i})', BLACK, fmt)
    if better == "lower":
        st = f'=IF(C{i}="","Open",IF(C{i}<=B{i},"On track","Behind"))'
    else:
        st = f'=IF(C{i}="","Open",IF(C{i}>=B{i},"On track","Behind"))'
    put(ws, f"E{i}", st)
    put(ws, f"F{i}", why).alignment = WRAP
    ws.row_dimensions[i].height = 30
ws.conditional_formatting.add("E8:E14", CellIsRule(operator="equal", formula=['"Behind"'], font=Font(name=FONT, color="C00000", bold=True)))
ws.conditional_formatting.add("E8:E14", CellIsRule(operator="equal", formula=['"On track"'], font=Font(name=FONT, color="008000", bold=True)))

put(ws, "A16", "Sunset list (fill per deal)", BOLD)
header(ws, 17, ["Item to retire", "Type", "Target date", "Done?", "", "Note"])
for i, (item, typ, note) in enumerate([
    ("Target's ERP / GL", "System", "Migrate to BSP chart of accounts at first close"),
    ("Target's billing/payment processors", "System", "Repoint payouts to Juno"),
    ("Target bank accounts", "Treasury", "Sweep, then close"),
    ("Redundant legal entities", "Entity", "Coordinate with tax (Italy transfer step)"),
    ("Target's board/management reporting packs", "Report", "Replace with BSP standard pack"),
], start=18):
    put(ws, f"A{i}", item, BLUE); put(ws, f"B{i}", typ, BLUE)
    put(ws, f"C{i}", None, BLUE, fill=YELLOW); put(ws, f"D{i}", None, BLUE, fill=YELLOW)
    put(ws, f"F{i}", note)

# ---------------------------------------------------------------- Lookback
ws = wb.create_sheet("Lookback")
style_ws(ws, {"A": 40, "B": 16, "C": 16, "D": 14, "E": 50})
put(ws, "A1", "12-month look-back: underwritten vs. actual", TITLE)
put(ws, "A2", "Run 12 months after close. The variances feed the next deal's underwriting assumptions. Example = Miro, underwritten values from the base case.", BLACK)
header(ws, 4, ["Assumption", "Underwritten", "Actual (12 mo)", "Variance", "What changes in the next underwriting"])
LB = [
    ("Revenue retained", "=Inputs!" + R["miro_ret"].split("!")[1], PCT),
    ("Operating margin", "=Inputs!" + R["miro_margin"].split("!")[1], PCT),
    ("Run-rate operating income ($mm)", "=Miro_Payback!B18", USD),
    ("Reorganization cost ($mm)", None, USD),
    ("Months to run-rate savings", None, '0'),
    ("Days to first consolidated close", "=Scorecard!B9", DAYS),
]
for i, (lab, uw, fmt) in enumerate(LB, start=5):
    put(ws, f"A{i}", lab)
    if uw:
        put(ws, f"B{i}", uw, GREEN, fmt)
    else:
        put(ws, f"B{i}", None, BLUE, fmt, fill=YELLOW)
    put(ws, f"C{i}", None, BLUE, fmt, fill=YELLOW)
    put(ws, f"D{i}", f'=IF(OR(B{i}="",C{i}=""),"",C{i}-B{i})', BLACK, fmt)
    put(ws, f"E{i}", None, BLUE, fill=YELLOW)
put(ws, "A12", "Tie-in: Bending Spoons' Minerva system predicts user lifetime value. The retained-revenue line should be checked against Minerva's cohort curves, not a separate spreadsheet.", BLACK).alignment = WRAP
ws.merge_cells("A12:E12")
ws.row_dimensions[12].height = 30

# ---------------------------------------------------------------- Sources
ws = wb.create_sheet("Sources")
style_ws(ws, {"A": 60, "B": 100})
put(ws, "A1", "Sources (opened Oct 12, 2026)", TITLE)
header(ws, 3, ["Source", "URL"])
for i, (name, url) in enumerate([
    ("Bending Spoons Q2 2026 earnings release (primary)", REL_URL),
    ("Bending Spoons: agreement to acquire Miro (primary)", MIRO_URL),
    ("Bending Spoons: completes acquisition of Airtable (primary)", AIRTABLE_URL),
    ("The Googly: prospectus analysis (secondary; covenant, debt rate, material weaknesses)", GOOGLY_URL),
    ("Mostly Metrics: S-1 breakdown (secondary; FY2025 revenue, NRR)", MM_URL),
], start=4):
    put(ws, f"A{i}", name)
    c = put(ws, f"B{i}", url)
    c.hyperlink = url

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT}")
